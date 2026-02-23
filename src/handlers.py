import os
import io
import time
import base64
import pandas as pd
from typing import List, Optional, Dict, Any
from docx import Document
from pdf2image import convert_from_path
from PIL import Image


class DocumentHandler:
    """Handles parsing of CSV/Excel, DOCX, and PDF documents for drug extraction."""

    def __init__(self, openai_client, deployment_name: str):
        self.client          = openai_client
        self.deployment_name = deployment_name

    def get_csv_sample(self, file_path: str) -> str:
        """
        Returns the first 5 rows from every sheet of an Excel file (or the single sheet of a CSV)
        as a labelled string, so the LLM can identify column structure per tab in one call.
        """
        try:
            ext = os.path.splitext(file_path)[1].lower()
            if ext not in (".xls", ".xlsx"):
                df = pd.read_csv(file_path, nrows=5, header=None)
                return f'SHEET "Sheet1":\n{df.to_string()}'

            xl     = pd.ExcelFile(file_path)
            parts  = []
            for sheet in xl.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet, nrows=5, header=None)
                parts.append(f'SHEET "{sheet}":\n{df.to_string()}')
            return "\n\n".join(parts)
        except Exception as e:
            return f"Error reading sample: {e}"

    def process_csv(self, file_path: str, drug_col_idx: int = 0, hcpcs_col_idx: Optional[int] = None, sheet_name: Optional[str] = None) -> List[str]:
        """Extracts drug name (and optionally HCPCS code) rows from a CSV or Excel file."""
        label = f"{os.path.basename(file_path)}" + (f" [{sheet_name}]" if sheet_name else "")
        print(f"   │    CSV    {label}")
        try:
            ext = os.path.splitext(file_path)[1].lower()
            if ext in (".xls", ".xlsx"):
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
            else:
                df = pd.read_csv(file_path, header=None)
            if df.empty:
                return []

            max_cols  = len(df.columns)
            drug_idx  = drug_col_idx if drug_col_idx < max_cols else 0
            skip_vals = {"nan", "null", "none", "drug name", "brand name", "name", "label"}

            extracted, found_content = [], False
            for index, row in df.iterrows():
                if row.isnull().all():
                    if found_content:
                        break   # Stop at first empty row after data has started
                    continue
                found_content = True

                drug_val = str(row.iloc[drug_idx]).strip()
                if not drug_val or drug_val.lower() in skip_vals:
                    continue

                entry = f"Drug: {drug_val}"
                if hcpcs_col_idx is not None and 0 <= hcpcs_col_idx < max_cols:
                    hcpcs_val = str(row.iloc[hcpcs_col_idx]).strip()
                    if hcpcs_val and hcpcs_val.lower() not in {"nan", "null", "none", "hcpcs", "j-code", "jcode"}:
                        entry += f" | Code: {hcpcs_val}"
                extracted.append(entry)

            return list(set(extracted))
        except Exception as e:
            print(f"   │    CSV    Error: {e}")
            return []

    def process_docx(self, file_path: str) -> str:
        """Extracts text from paragraphs and tables in a DOCX file."""
        print(f"   │   DOCX    {os.path.basename(file_path)}")
        try:
            doc   = Document(file_path)
            lines = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
        except Exception as e:
            print(f"   │   DOCX    Error: {e}")
            return ""

    def process_pdf_vision(self, file_path: str, stats: Optional[Dict[str, Any]] = None) -> str:
        """Converts first 5 PDF pages to images and runs a single batched Vision call."""
        print(f"   │    PDF    {os.path.basename(file_path)} (vision — batch)")
        try:
            images = convert_from_path(file_path, first_page=1, last_page=5)
            os.makedirs("cache", exist_ok=True)

            page_images_b64 = []
            for i, image in enumerate(images):
                image.save(f"cache/{os.path.basename(file_path)}_page_{i+1}.png", "PNG")
                buf = io.BytesIO()
                image.save(buf, format="PNG")
                page_images_b64.append(base64.b64encode(buf.getvalue()).decode("utf-8"))

            print(f"   │    PDF    {len(page_images_b64)} page(s) → 1 vision call")
            return self._call_vision_model_batch(page_images_b64, stats=stats)
        except Exception as e:
            print(f"   │    PDF    Error: {e}")
            return ""

    def _call_vision_model_batch(self, base64_images: List[str], stats: Optional[Dict[str, Any]] = None) -> str:
        """Sends all page images in a single Vision API call and returns the combined text."""
        try:
            # Build content: one text instruction + one image block per page
            content = [
                {
                    "type": "text", 
                    "text": "Analyze these pages from a medical policy document. Extract all drug names and HCPCS/J-codes, especially from exclusion or carve-out lists. Combine results across all pages."
                }
            ]
            for img_b64 in base64_images:
                content.append({
                    "type": "image_url", 
                    "image_url": {"url": f"data:image/png;base64,{img_b64}"}
                })

            t0       = time.time()
            response = self.client.chat.completions.create(
                model=self.deployment_name,
                messages=[
                    {
                        "role": "system", 
                        "content": "You are a specialized medical document analyzer. Extract drug names, HCPCS codes, and exclusion lists from the provided pages."
                    },
                    {
                        "role": "user", 
                        "content": content
                    },
                ],
                max_completion_tokens=2000,
            )
            elapsed = time.time() - t0
            if stats is not None and hasattr(response, "usage") and response.usage:
                stats["llm_calls"]        += 1
                stats["input_tokens"]     += response.usage.prompt_tokens
                stats["output_tokens"]    += response.usage.completion_tokens
                stats["total_tokens"]     += response.usage.total_tokens
                stats["duration_seconds"]  = round(stats["duration_seconds"] + elapsed, 3)
            return response.choices[0].message.content or ""
        except Exception as e:
            print(f"   │  Vision   Error: {e}")
            return ""